"""Creates data/worksheet_database.xlsx with 5 sample rows (+ README sheet)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEAD = ["ID", "Category", "Object_EN", "Object_ID", "Relation_EN", "Relation_ID", "Image_Prompt"]
ROWS = [
    ["001", "Animals & Their Products", "Chicken", "Ayam", "Egg", "Telur", "cute chicken holding egg"],
    ["002", "Animals & Their Products", "Bee", "Lebah", "Honey", "Madu", "cute bee holding honey jar"],
    ["003", "Animals & Their Products", "Cow", "Sapi", "Milk", "Susu", "cute cow with milk carton"],
    ["004", "Animals & Their Products", "Sheep", "Domba", "Wool", "Wol", "cute sheep with ball of wool"],
    ["005", "Plants & Their Fruits", "Tree", "Pohon", "Apple", "Apel", "cute tree with red apples"],
]
wb = Workbook(); ws = wb.active; ws.title = "Database"
ws.append(HEAD)
for r in ROWS:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4FB3E8")
for col, w in zip("ABCDEFG", [8, 28, 14, 14, 14, 14, 44]):
    ws.column_dimensions[col].width = w
for r in range(2, 2000):
    ws.cell(r, 1).number_format = "@"          # keep 001 as text
ws.freeze_panes = "A2"
dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="14",
                    showErrorMessage=True, error="Keep words short (max 14 letters) for ages 3-5")
ws.add_data_validation(dv); dv.add("C2:F2000")

# ---- optional sheet: bundles (every page may use a DIFFERENT row)
bs = wb.create_sheet("Bundles")
bs.append(["Bundle", "Template", "ID"])
for i, rid in enumerate(["001", "002", "003", "004", "005", "001", "002", "003", "004", "005"]):
    bs.append(["Animal_Friends_Pack_1", "", rid])                 # Template kosong = urut A, B, C ... J
for t, rid in [("A", "002"), ("C", "004"), ("J", "005")]:
    bs.append(["Mini_Pack", t, rid])                              # Template diisi = halaman tertentu
for c in bs[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="9B7BE0")
for col, w in zip("ABC", [28, 12, 8]):
    bs.column_dimensions[col].width = w
for r in range(2, 5000):
    bs.cell(r, 3).number_format = "@"

rd = wb.create_sheet("README")
for line in [
    "Columns A-G are required. Optional extra columns (add in H+): Countable (Y/N), Templates (e.g. A,B,C), Status.",
    "ID: unique text, e.g. 001.  Category: becomes the folder name.  Image_Prompt: must mention Object_EN.",
    "Words must exist in data/approved_vocab.csv (verified translations) or QC will warn / fail.",
    "Sheet 'Bundles' (optional): Bundle = pack name, Template = A..J (blank = A, B, C... in order), ID = row of sheet Database. One line = one page.",
]:
    rd.append([line])
rd.column_dimensions["A"].width = 120
wb.save("data/worksheet_database.xlsx"); print("saved data/worksheet_database.xlsx")
